import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_FONT = "Arial"
ACCENT_BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "6B7280"
TABLE_HEADER = "F2F4F7"
CODE_FILL = "F4F6F9"
BORDER = "D9E2EC"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, name=BASE_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    rfonts = run._element.rPr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        run._element.rPr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), name)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    existing_grid = tbl.find(qn("w:tblGrid"))
    if existing_grid is not None:
        tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_borders(cell)


def set_repeating_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def next_numbering_ids(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
        if element.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
        if element.get(qn("w:numId")) is not None
    ]
    return (max(abstract_ids, default=0) + 1, max(num_ids, default=0) + 1)


def create_numbered_list(doc):
    numbering = doc.part.numbering_part.element
    abstract_id, num_id = next_numbering_ids(doc)

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))

    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)

    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)

    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    level.append(lvl_text)

    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)

    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)

    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")

    num_id_element = num_pr.find(qn("w:numId"))
    if num_id_element is None:
        num_id_element = OxmlElement("w:numId")
        num_pr.append(num_id_element)
    num_id_element.set(qn("w:val"), str(num_id))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("หน้า ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, ACCENT_BLUE, 16, 8),
        ("Heading 2", 13, ACCENT_BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = BASE_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = BASE_FONT
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.text = ""
    left = footer_para.add_run("SRS Portfolio System")
    set_run_font(left, size=9, color=MUTED)
    footer_para.add_run("\t")
    add_page_number(footer_para)


def add_inline_markdown(paragraph, text, base_size=11):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=base_size, color=INK)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=INK)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Courier New", size=max(base_size - 1, 8), color="374151")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, color=INK)


def clean_text(text):
    return text.strip().replace("  ", " ")


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def compute_widths(rows):
    col_count = max(len(row) for row in rows)
    if col_count == 1:
        return [TABLE_WIDTH_DXA]
    if col_count == 2:
        first_header = rows[0][0].lower()
        if first_header in {"field", "token", "method"} or "field" in first_header:
            return [2500, 6860]
        return [2800, 6560]
    if col_count == 3:
        return [1900, 2550, 4910]
    if col_count == 4:
        return [1850, 2550, 2450, 2510]
    if col_count == 5:
        return [2100, 1750, 1750, 1750, 2010]

    weights = []
    for idx in range(col_count):
        max_len = max((len(row[idx]) if idx < len(row) else 0) for row in rows)
        weights.append(max(8, min(max_len, 36)))
    total = sum(weights)
    raw = [int(TABLE_WIDTH_DXA * weight / total) for weight in weights]
    raw[-1] += TABLE_WIDTH_DXA - sum(raw)
    return raw


def add_markdown_table(doc, table_rows):
    rows = [row for row in table_rows if row and not is_separator_row(row)]
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    widths = compute_widths(rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeating_header(table.rows[0])

    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_shading(cell, TABLE_HEADER)
            text = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.text = ""
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_markdown(paragraph, text, base_size=9.5)
            for run in paragraph.runs:
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("111827")

    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=6)


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    para = cell.paragraphs[0]
    para.text = ""
    para.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(code.splitlines()):
        if idx:
            para.add_run().add_break()
        run = para.add_run(line)
        set_run_font(run, name="Courier New", size=8.5, color="374151")
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=6)


def add_heading(doc, level, text):
    style = f"Heading {min(level, 3)}"
    paragraph = doc.add_paragraph(style=style)
    if level >= 4:
        paragraph.style = doc.styles["Heading 3"]
    paragraph.text = ""
    run = paragraph.add_run(text)
    if level >= 4:
        set_run_font(run, size=11, bold=True, color=DARK_BLUE)
    else:
        # Style carries the preset; explicit font helps Thai fallback.
        set_run_font(
            run,
            size={1: 16, 2: 13, 3: 12}.get(level, 12),
            bold=True,
            color=ACCENT_BLUE if level <= 2 else DARK_BLUE,
        )
    paragraph.paragraph_format.keep_with_next = True


def add_title_block(doc, title, subtitle, metadata):
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_para.paragraph_format.space_before = Pt(0)
    title_para.paragraph_format.space_after = Pt(4)
    run = title_para.add_run(title)
    set_run_font(run, size=22, bold=True, color=ACCENT_BLUE)

    subtitle_para = doc.add_paragraph()
    subtitle_para.paragraph_format.space_after = Pt(12)
    run = subtitle_para.add_run(subtitle)
    set_run_font(run, size=14, bold=True, color=DARK_BLUE)

    if metadata:
        table = doc.add_table(rows=len(metadata), cols=2)
        table.style = "Table Grid"
        set_table_geometry(table, [2600, 6760])
        for idx, (label, value) in enumerate(metadata):
            label_cell = table.cell(idx, 0)
            value_cell = table.cell(idx, 1)
            set_cell_shading(label_cell, TABLE_HEADER)
            for cell in (label_cell, value_cell):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                para = cell.paragraphs[0]
                para.text = ""
                para.paragraph_format.space_after = Pt(0)
            label_run = label_cell.paragraphs[0].add_run(label)
            set_run_font(label_run, size=10, bold=True, color="111827")
            value_run = value_cell.paragraphs[0].add_run(value)
            set_run_font(value_run, size=10, color=INK)

    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=12)


def parse_front_matter(lines):
    idx = 0
    while idx < len(lines) and not lines[idx].strip():
        idx += 1
    title = re.sub(r"^#\s+", "", lines[idx].strip()) if idx < len(lines) else "SRS"
    idx += 1
    while idx < len(lines) and not lines[idx].strip():
        idx += 1
    subtitle = re.sub(r"^##\s+", "", lines[idx].strip()) if idx < len(lines) else ""
    idx += 1
    metadata = []
    meta_pattern = re.compile(r"^\*\*(.+?):\*\*\s*(.+?)\s*$")
    while idx < len(lines):
        line = lines[idx].strip()
        if line == "---":
            idx += 1
            break
        if line:
            match = meta_pattern.match(line)
            if match:
                metadata.append((match.group(1), match.group(2)))
        idx += 1
    return title, subtitle, metadata, idx


def build_docx(markdown_path, output_path):
    source = Path(markdown_path)
    lines = source.read_text(encoding="utf-8").splitlines()

    doc = Document()
    configure_document(doc)

    title, subtitle, metadata, idx = parse_front_matter(lines)
    add_title_block(doc, title, subtitle, metadata)

    in_code = False
    code_lines = []
    table_buffer = []
    current_num_id = None

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

    while idx < len(lines):
        raw_line = lines[idx]
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                flush_table()
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, "\n".join(code_lines))
                in_code = False
                code_lines = []
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if not stripped:
            flush_table()
            current_num_id = None
            idx += 1
            continue

        if stripped == "---":
            flush_table()
            current_num_id = None
            spacer = doc.add_paragraph()
            set_paragraph_spacing(spacer, before=0, after=4)
            idx += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_buffer.append(split_table_row(stripped))
            idx += 1
            continue

        flush_table()

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            current_num_id = None
            level = len(heading_match.group(1))
            add_heading(doc, level, clean_text(heading_match.group(2)))
            idx += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            current_num_id = None
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            add_inline_markdown(paragraph, clean_text(bullet_match.group(1)))
            idx += 1
            continue

        numbered_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered_match:
            if current_num_id is None:
                current_num_id = create_numbered_list(doc)
            paragraph = doc.add_paragraph()
            set_paragraph_spacing(paragraph, before=0, after=8, line=1.167)
            apply_numbering(paragraph, current_num_id)
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            add_inline_markdown(paragraph, clean_text(numbered_match.group(1)))
            idx += 1
            continue

        current_num_id = None
        paragraph = doc.add_paragraph()
        set_paragraph_spacing(paragraph, before=0, after=6, line=1.10)
        add_inline_markdown(paragraph, clean_text(stripped))
        idx += 1

    flush_table()
    if in_code and code_lines:
        add_code_block(doc, "\n".join(code_lines))

    doc.core_properties.title = "Software Requirements Specification - Portfolio System"
    doc.core_properties.subject = "ระบบ Portfolio ออนไลน์สำหรับนักศึกษาวิศวกรรมคอมพิวเตอร์"
    doc.core_properties.comments = "Generated from SRS-Portfolio-System.md"
    doc.save(output_path)


if __name__ == "__main__":
    build_docx("SRS-Portfolio-System.md", "SRS-Portfolio-System.docx")
